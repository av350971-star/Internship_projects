import io
import re
import json
import hashlib
import uuid
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional
from datetime import datetime

from backend.config import settings
from backend.storage.sqlite_store import store

class DocumentIngestionService:
    def __init__(self):
        self.chunk_size = settings.CHUNK_SIZE
        self.chunk_overlap = settings.CHUNK_OVERLAP

    @staticmethod
    def calculate_checksum(data: bytes) -> str:
        """Calculates SHA-256 checksum of raw file bytes."""
        return hashlib.sha256(data).hexdigest()

    def run_ocr_on_image(self, img_bytes: bytes) -> str:
        """Runs Tesseract OCR on raw image bytes."""
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(io.BytesIO(img_bytes))
            text = pytesseract.image_to_string(img).strip()
            return text
        except Exception as e:
            print(f"OCR processing error: {e}")
            return ""

    def extract_text_from_pdf(self, file_bytes: bytes) -> List[Tuple[int, str]]:
        """
        Extracts text from PDF page by page.
        If a page contains no digital text (e.g. scanned image / photo document),
        it automatically triggers OCR using Tesseract!
        """
        pages = []
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text").strip()

                # If page has very little or no digital text, run OCR fallback!
                if len(text) < 20:
                    try:
                        pix = page.get_pixmap(dpi=150)
                        ocr_result = self.run_ocr_on_image(pix.tobytes("png"))
                        if ocr_result and len(ocr_result) > len(text):
                            text = f"[OCR Extracted]\n{ocr_result}"
                    except Exception as ocr_err:
                        print(f"PDF Page {page_num+1} OCR fallback error: {ocr_err}")

                if text:
                    pages.append((page_num + 1, text))
            doc.close()
            if pages:
                return pages
        except Exception as e:
            print(f"PyMuPDF error: {e}")

        # Fallback to pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages.append((page_num + 1, text.strip()))
        except Exception as e:
            print(f"pypdf error: {e}")

        return pages

    def extract_text_from_image(self, file_bytes: bytes) -> List[Tuple[int, str]]:
        """Extracts text from standalone images (PNG, JPG, JPEG) using OCR."""
        text = self.run_ocr_on_image(file_bytes)
        if text:
            return [(1, f"[OCR Image Content]\n{text}")]
        return [(1, "Scanned Image (No readable text recognized via OCR)")]

    def extract_text_from_docx(self, file_bytes: bytes) -> List[Tuple[int, str]]:
        """Extracts text from DOCX file."""
        pages = []
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        full_text.append(row_text)
            text = "\n\n".join(full_text)
            if text:
                pages.append((1, text))
        except Exception as e:
            print(f"Error extracting DOCX: {e}")
        return pages

    def extract_text_from_json(self, file_bytes: bytes) -> List[Tuple[int, str]]:
        """Extracts text from structured JSON, converting objects to readable key-value text."""
        try:
            raw_text = file_bytes.decode("utf-8", errors="replace")
            parsed = json.loads(raw_text)
            lines = []
            if isinstance(parsed, list):
                for i, item in enumerate(parsed):
                    lines.append(f"--- Item {i+1} ---")
                    if isinstance(item, dict):
                        for k, v in item.items():
                            lines.append(f"{k}: {v}")
                    else:
                        lines.append(str(item))
            elif isinstance(parsed, dict):
                for k, v in parsed.items():
                    if isinstance(v, (dict, list)):
                        lines.append(f"{k}: {json.dumps(v, indent=2)}")
                    else:
                        lines.append(f"{k}: {v}")
            else:
                lines.append(str(parsed))
            return [(1, "\n".join(lines))]
        except Exception as e:
            print(f"Error extracting JSON: {e}")
            return [(1, file_bytes.decode("utf-8", errors="replace"))]

    def extract_text_from_plain(self, file_bytes: bytes) -> List[Tuple[int, str]]:
        """Extracts text from TXT or MD files."""
        try:
            text = file_bytes.decode("utf-8", errors="replace").strip()
            return [(1, text)] if text else []
        except Exception as e:
            print(f"Error extracting plain text: {e}")
            return []

    def extract_text(self, file_bytes: bytes, filename: str) -> List[Tuple[int, str]]:
        """Extracts text pages based on file extension."""
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return self.extract_text_from_pdf(file_bytes)
        elif ext in [".docx", ".doc"]:
            return self.extract_text_from_docx(file_bytes)
        elif ext == ".json":
            return self.extract_text_from_json(file_bytes)
        elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"]:
            return self.extract_text_from_image(file_bytes)
        elif ext in [".txt", ".md", ".markdown", ".csv", ".log"]:
            return self.extract_text_from_plain(file_bytes)
        else:
            # Fallback to UTF-8 decoding
            return self.extract_text_from_plain(file_bytes)

    def chunk_page_text(self, text: str, page_number: int, doc_id: str, tenant_id: str, start_index: int = 0) -> List[Dict[str, Any]]:
        """
        Splits page text into clean chunks with overlap, preserving page and section metadata.
        """
        chunks = []
        if not text:
            return chunks

        # Normalize whitespace while preserving paragraphs
        paragraphs = re.split(r'\n{2,}', text)
        clean_text = "\n\n".join(p.strip() for p in paragraphs if p.strip())
        
        # Sliding window chunker
        start = 0
        text_len = len(clean_text)
        current_idx = start_index

        while start < text_len:
            end = min(start + self.chunk_size, text_len)
            
            # If we are in the middle of a sentence, try finding a sentence boundary
            if end < text_len:
                boundary = clean_text.rfind(". ", start, end)
                if boundary != -1 and boundary > start + (self.chunk_size // 2):
                    end = boundary + 1

            chunk_content = clean_text[start:end].strip()
            if len(chunk_content) > 20:  # Avoid empty/tiny chunks
                # Detect heading / section title if first line looks like one
                first_line = chunk_content.split("\n")[0][:60]
                section_title = first_line if first_line.startswith("#") or ":" in first_line else ""

                chunk_id = f"chk_{doc_id}_{current_idx}"
                chunks.append({
                    "id": chunk_id,
                    "doc_id": doc_id,
                    "tenant_id": tenant_id,
                    "chunk_index": current_idx,
                    "page_number": page_number,
                    "section_title": section_title,
                    "text_content": chunk_content,
                    "token_count": len(chunk_content.split())
                })
                current_idx += 1

            if end >= text_len:
                break

            # Advance window ensuring proper step forward
            advance = max((end - start) - self.chunk_overlap, 100)
            start += advance

        return chunks

    def process_and_ingest(
        self,
        file_bytes: bytes,
        filename: str,
        tenant_id: str,
        project_id: str = "default",
        force_reindex: bool = False,
        embedding_service: Any = None
    ) -> Dict[str, Any]:
        """
        Full ingestion pipeline:
        1. Calculates SHA-256 checksum
        2. Detects unchanged duplicates vs updates vs new files
        3. Parses multi-format content
        4. Chunks text with metadata
        5. Computes embeddings
        6. Persists to SQLite
        """
        checksum = self.calculate_checksum(file_bytes)
        file_size = len(file_bytes)
        ext = Path(filename).suffix.lower()

        # Check if identical file already exists for this tenant
        existing_by_hash = store.find_document_by_checksum(tenant_id, checksum)
        if existing_by_hash and not force_reindex:
            return {
                "status": "unchanged",
                "message": f"Document '{filename}' already ingested with identical content.",
                "doc_id": existing_by_hash["id"],
                "filename": filename,
                "tenant_id": tenant_id,
                "chunk_count": existing_by_hash["chunk_count"],
                "checksum": checksum
            }

        # Check if existing document by filename exists (Update/Re-index flow)
        existing_by_name = store.find_document_by_filename(tenant_id, filename)
        if existing_by_name:
            doc_id = existing_by_name["id"]
            action = "reindexed"
            # Delete old chunks for clean update
            store.delete_chunks_for_document(doc_id)
        else:
            doc_id = f"doc_{uuid.uuid4().hex[:10]}"
            action = "indexed"

        # Extract text pages
        pages = self.extract_text(file_bytes, filename)
        if not pages:
            # Fallback if binary extraction yielded no pages
            pages = [(1, file_bytes.decode("utf-8", errors="ignore"))]

        # Generate chunks
        all_chunks = []
        chunk_idx = 0
        for page_num, page_text in pages:
            page_chunks = self.chunk_page_text(page_text, page_num, doc_id, tenant_id, chunk_idx)
            chunk_idx += len(page_chunks)
            all_chunks.extend(page_chunks)

        if not all_chunks:
            # Empty file fallback
            all_chunks.append({
                "id": f"chk_{doc_id}_0",
                "doc_id": doc_id,
                "tenant_id": tenant_id,
                "chunk_index": 0,
                "page_number": 1,
                "section_title": "",
                "text_content": f"Document: {filename} (No readable text extracted)",
                "token_count": 6
            })

        # Generate embeddings if service provided
        if embedding_service:
            texts = [c["text_content"] for c in all_chunks]
            embeddings = embedding_service.encode_batch(texts)
            for i, emb in enumerate(embeddings):
                all_chunks[i]["embedding"] = emb
        else:
            for c in all_chunks:
                c["embedding"] = None

        # Persist raw file to data/uploaded_files for future complete re-indexing
        try:
            upload_dir = settings.DATA_DIR / "uploaded_files" / tenant_id
            upload_dir.mkdir(parents=True, exist_ok=True)
            saved_file_path = upload_dir / f"{doc_id}_{filename}"
            with open(saved_file_path, "wb") as f:
                f.write(file_bytes)
        except Exception as e:
            print(f"Warning: Could not save raw file to disk: {e}")

        # Save document record
        doc_data = {
            "id": doc_id,
            "filename": filename,
            "file_type": ext or "txt",
            "file_size": file_size,
            "tenant_id": tenant_id,
            "project_id": project_id,
            "checksum": checksum,
            "chunk_count": len(all_chunks)
        }
        store.save_document(doc_data)

        # Save chunks in SQLite
        store.save_chunks(all_chunks)

        return {
            "status": action,
            "message": f"Successfully {action} '{filename}' with {len(all_chunks)} chunks.",
            "doc_id": doc_id,
            "filename": filename,
            "tenant_id": tenant_id,
            "project_id": project_id,
            "chunk_count": len(all_chunks),
            "checksum": checksum
        }

ingestion_service = DocumentIngestionService()
